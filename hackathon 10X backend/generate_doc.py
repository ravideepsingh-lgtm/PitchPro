"""
Script to generate a Word document containing all agent prompts and data keys.
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

doc = Document()

# ── Title ──
title = doc.add_heading("Agentic Batch Upsell Pipeline", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("Complete Prompt & Data Key Reference Document")
doc.add_paragraph(f"Project: BJP (Beyond Just Product) - 10X Productivity Backend")
doc.add_paragraph("")

# ═══════════════════════════════════════════════════
# SECTION 1: Pipeline Architecture
# ═══════════════════════════════════════════════════
doc.add_heading("1. Pipeline Architecture", level=1)
doc.add_paragraph(
    "The system uses a 5-agent, 3-phase pipeline to analyze enterprise seller data "
    "and produce a final upsell recommendation for each seller."
)

doc.add_heading("Execution Flow", level=2)
table = doc.add_table(rows=4, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Phase"
hdr[1].text = "Agents"
hdr[2].text = "Execution"
row1 = table.rows[1].cells
row1[0].text = "Phase 1"
row1[1].text = "Profile Agent + Positive Agent + Negative Agent"
row1[2].text = "Parallel (3 simultaneous LLM calls)"
row2 = table.rows[2].cells
row2[0].text = "Phase 2"
row2[1].text = "Sentiment Agent"
row2[2].text = "Sequential (receives Phase 1 outputs)"
row3 = table.rows[3].cells
row3[0].text = "Phase 3"
row3[1].text = "Final Upsell Decision Agent"
row3[2].text = "Sequential (receives all 4 agent outputs)"

doc.add_paragraph("")

# ═══════════════════════════════════════════════════
# SECTION 2: Data Sources
# ═══════════════════════════════════════════════════
doc.add_heading("2. Data Sources (Excel Sheets)", level=1)
doc.add_paragraph(
    "The system loads data from a master .xlsb workbook (Paid Clients) and multiple standalone CSV/Excel files."
)

data_sources = [
    ("Client Base", "Master File (Paid Clients)", "custtype_name, company, city, catalog_service, ive_service, service_start_date, service_end_date, enabled, STATE, client_since, service_mode, pin, gst, ANNUAL_TURNOVER, Company Type, location_pref, txn_30_days, txn_14_days, txn_7_days, active_days_30/14/7, client_vintage"),
    ("BL Allocation", "Master File (Paid Clients)", "fk_glusr_usr_id, weekly, hosting, daily"),
    ("DSR Summary", "Master File (Paid Clients)", "fk_glusr_usr_id, summary, insertion_date, status, updated_on"),
    ("BL Data 1", "Master File (Paid Clients)", "seller_id, custtype_name, city_name, state_name, location_pref, txn_30_days, txn_14_days, txn_7_days, active_days_30/14/7"),
    ("MCAT Rank", "Master File (Paid Clients)", "seller_id, custtype_name, rank_a_mcat_count, rank_b_mcat_count, rank_c/d/f_mcat_count, high/medium/low_sold_mcat_count"),
    ("Upsell", "Master File (Paid Clients)", "GluserId, Segment Recommendation"),
    ("Scorecard", "Master File (Paid Clients)", "fk_glusr_usr_id, ast_buy_enq, avg_ratings, connect_duration, count_of_number_mapped, highest_service"),
    ("IMA", "Master File (Paid Clients)", "fk_glusr_usr_id, weekly, hosting, daily"),
    ("csd_kcd_active_seller", "Standalone CSV", "glusr_usr_id, ownership_type, me_flag, sales_vertical_name, tele_vertical_name, totalcalls_30/60/90, total_ans_calls_30/60/90, prime_final_rag_score, industry, flag_gst/tan/tin/iec, count_mcats, creationdate, last_catalog_update_date, last_call_date, dsr_meeting_last_date"),
    ("ima_cron_daily", "Standalone CSV", "glusr_usr_id, bl_cred_indian, bl_cred_foreign_daily/weekly, total_used, total_active, total_bl_active, total_cv_active, in/fl_total_active, in/fl_bl_active, total_cons_indian/foreign, bl/cv_cons_indian/foreign, gst_turnover"),
    ("pns_csd_daily_cron", "Standalone CSV", "glusr_usr_id, pns_call, success_call, total_enq, total_astbuy, callbacks_total/indian/foreign, lms_total/indian/foreign, seller_rating"),
    ("customer services 1st may", "Standalone CSV/XLSB", "glid, service_name, service_id, st_dt, end_dt, st_dt_renewal_enddate, status, monthly_type, ptype, ar_invoice_total_amount, ar_invoice_os_amount, ar_recpt_status"),
    ("seller_mcat data", "Standalone CSV", "seller_id, custtype_name, mcat_id, mcat_name, rank, sold_type"),
    ("NI", "Standalone CSV", "seller_id, total_ni, wrong_category_ni, location_ni, retail_ni, other_ni"),
    ("QRF", "Standalone CSV", "seller_id, total_irr_qrf_count, wrong_category, location_mismatch, retail_qrf, other_qrf"),
    ("Notif", "Standalone CSV", "fk_glusr_usr_id, notif_trig, notif_sent, notif_del, notif_open, notif_swiped, uniq_seller_notif_open"),
]

table2 = doc.add_table(rows=len(data_sources) + 1, cols=3)
table2.style = "Light Grid Accent 1"
hdr2 = table2.rows[0].cells
hdr2[0].text = "Sheet/File Name"
hdr2[1].text = "Source"
hdr2[2].text = "Key Columns"
for i, (name, src, cols) in enumerate(data_sources):
    row = table2.rows[i + 1].cells
    row[0].text = name
    row[1].text = src
    row[2].text = cols

doc.add_page_break()

# ═══════════════════════════════════════════════════
# SECTION 3: AGENT 1 - Profile Agent
# ═══════════════════════════════════════════════════
doc.add_heading("3. Agent 1: Profile Agent (Seller Profiling)", level=1)

doc.add_heading("Purpose", level=2)
doc.add_paragraph("Creates a comprehensive business/customer profile of the seller without making upsell decisions.")

doc.add_heading("Input Data Keys", level=2)
profile_keys = [
    ("customer_profile", "Client Base + csd_kcd_active_seller + ima_cron_daily", 
     "custtype_name, ownership_type, me_flag, client_since, client_vintage, creationdate, industry, Company Type, ANNUAL_TURNOVER, gst_turnover"),
    ("location_profile", "Client Base", "city, STATE, pin, location_pref"),
    ("service_profile", "Client Base + Customer Services 1st May", 
     "catalog_service, ive_service, service_mode, enabled, service_name, service_id, status, st_dt, end_dt, st_dt_renewal_enddate, monthly_type, ptype, ar_invoice_total_amount, ar_invoice_os_amount, ar_recpt_status"),
    ("trust_profile", "csd_kcd_active_seller", "flag_gst, flag_tan, flag_tin, flag_iec"),
    ("category_profile", "csd_kcd_active_seller + MCAT Rank + seller_mcat data", 
     "count_mcats, rank_a_mcat_count, rank_b_mcat_count, high_sold_mcat_count, medium_sold_mcat_count, mcat_id, mcat_name, rank, sold_type"),
    ("activity_profile", "BL Data 1 + csd_kcd_active_seller", 
     "txn_30_days, txn_14_days, txn_7_days, active_days_30/14/7, last_catalog_update_date, last_call_date, dsr_meeting_last_date"),
    ("communication_profile", "csd_kcd_active_seller + pns_csd_daily_cron", 
     "totalcalls_30, totalcalls_60, total_ans_calls_30, pns_call, success_call, seller_rating"),
    ("internal_business_context", "csd_kcd_active_seller + Scorecard", 
     "sales_vertical_name, prime_final_rag_score, highest_service"),
    ("latest_dsr_summary", "DSR Summary", "summary, status"),
]

table3 = doc.add_table(rows=len(profile_keys) + 1, cols=3)
table3.style = "Light Grid Accent 1"
hdr3 = table3.rows[0].cells
hdr3[0].text = "Section"
hdr3[1].text = "Source Sheet(s)"
hdr3[2].text = "Keys Used"
for i, (section, source, keys) in enumerate(profile_keys):
    row = table3.rows[i + 1].cells
    row[0].text = section
    row[1].text = source
    row[2].text = keys

doc.add_heading("Full Prompt", level=2)
profile_prompt = """You are a Seller Profiling Agent.

Your job is only to create a business/customer profile of the seller.

Do NOT decide whether to upsell.
Do NOT recommend any service.
Do NOT mention company name or personal identity.
Do NOT invent missing information.
Use "Unknown" when data is missing or unclear.

Analyze the structured seller data and return valid JSON only.

You must generate:
1. seller_profile_type
2. business_maturity
3. client_vintage
4. current_customer_type
5. service_status
6. business_category_strength
7. catalog_completeness
8. engagement_level
9. location_profile
10. trust_profile
11. profile_risks
12. profile_summary

Use evidence from fields wherever possible.

Seller Data:
{seller_feature_json}

Return JSON in this schema:
{
  "seller_profile_type": "",
  "business_maturity": "High / Medium / Low / Unknown",
  "client_vintage": "Long-term / Medium / New / Unknown",
  "current_customer_type": "",
  "service_status": "Active / Expired / Renewal Due / Disabled / Unknown",
  "business_category_strength": "Strong / Moderate / Weak / Unknown",
  "catalog_completeness": "Good / Average / Poor / Unknown",
  "engagement_level": "High / Medium / Low / Unknown",
  "location_profile": {
    "city": "",
    "state": "",
    "location_preference": ""
  },
  "trust_profile": {
    "gst_available": false,
    "tan_available": false,
    "tin_available": false,
    "iec_available": false
  },
  "profile_risks": [],
  "evidence_fields": [],
  "profile_summary": ""
}"""
doc.add_paragraph(profile_prompt, style="No Spacing")

doc.add_page_break()

# ═══════════════════════════════════════════════════
# SECTION 4: AGENT 2 - Positive Engagement Agent
# ═══════════════════════════════════════════════════
doc.add_heading("4. Agent 2: Positive Engagement Agent", level=1)

doc.add_heading("Purpose", level=2)
doc.add_paragraph("Identifies all positive signals indicating the seller is actively using and benefiting from the platform.")

doc.add_heading("Input Data Keys", level=2)
pos_keys = [
    ("bl_consumption", "ima_cron_daily", "total_used, total_active, total_bl_active, total_cv_active, in_total_active, in_bl_active, fl_total_active, fl_bl_active, total_cons_indian, bl_cons_indian, cv_cons_indian, total_cons_foreign, bl_cons_foreign, cv_cons_foreign, bl_cred_indian, bl_cred_foreign_daily, bl_cred_foreign_weekly"),
    ("lms_engagement", "pns_csd_daily_cron", "lms_total, lms_indian, lms_foreign"),
    ("enquiry_engagement", "pns_csd_daily_cron + Scorecard", "total_enq, total_astbuy, ast_buy_enq, callbacks_total, callbacks_indian, callbacks_foreign"),
    ("bl_activity", "BL Data 1", "txn_30_days, txn_14_days, txn_7_days, active_days_30, active_days_14, active_days_7"),
    ("notification_engagement", "Notif", "notif_trig, notif_sent, notif_del, notif_open, notif_swiped, uniq_seller_notif_open"),
    ("service_context", "Client Base + Scorecard + pns_csd_daily_cron", "catalog_service, highest_service, seller_rating, avg_ratings, connect_duration"),
]
table4 = doc.add_table(rows=len(pos_keys) + 1, cols=3)
table4.style = "Light Grid Accent 1"
hdr4 = table4.rows[0].cells
hdr4[0].text = "Section"
hdr4[1].text = "Source Sheet(s)"
hdr4[2].text = "Keys Used"
for i, (section, source, keys) in enumerate(pos_keys):
    row = table4.rows[i + 1].cells
    row[0].text = section
    row[1].text = source
    row[2].text = keys

doc.add_heading("Full Prompt", level=2)
pos_prompt = """You are a Positive Engagement Analysis Agent for an enterprise seller platform.

Your job is to analyze the seller's engagement data and identify ALL positive signals that indicate the seller is actively using the platform and benefiting from it.

Do NOT decide whether to upsell.
Do NOT recommend any service.
Do NOT mention company name or personal identity.
Do NOT invent missing information.
Use "Unknown" when data is missing or unclear.

Focus on these signal categories:
1. BL Consumption: How actively is the seller consuming buy leads (Indian + Foreign)?
2. LMS Engagement: Are they using Lead Management System?
3. Enquiry & AstBuy: Volume of enquiries and assisted buys.
4. BL Activity: Transaction volume and active days in last 30/14/7 days.
5. Notification Engagement: Are they opening and engaging with notifications?
6. Service Context: Current service level, ratings, and connect duration.

Seller Positive Engagement Data:
{seller_feature_json}

Return JSON in this schema:
{
  "bl_consumption_level": "High / Medium / Low / None",
  "lms_engagement_level": "High / Medium / Low / None",
  "enquiry_responsiveness": "High / Medium / Low / None",
  "bl_activity_level": "High / Medium / Low / None",
  "notification_engagement": "High / Medium / Low / None",
  "overall_positive_score": "Strong / Moderate / Weak / None",
  "positive_signals": [],
  "evidence_fields": [],
  "positive_summary": ""
}"""
doc.add_paragraph(pos_prompt, style="No Spacing")

doc.add_page_break()

# ═══════════════════════════════════════════════════
# SECTION 5: AGENT 3 - Negative Engagement Agent
# ═══════════════════════════════════════════════════
doc.add_heading("5. Agent 3: Negative Engagement Agent", level=1)

doc.add_heading("Purpose", level=2)
doc.add_paragraph("Identifies all negative signals, risks, and friction points indicating dissatisfaction, churn risk, or poor platform experience.")

doc.add_heading("Input Data Keys", level=2)
neg_keys = [
    ("ni_signals", "NI", "total_ni, wrong_category_ni, location_ni, retail_ni, other_ni"),
    ("qrf_signals", "QRF", "total_irr_qrf_count, wrong_category, location_mismatch, retail_qrf, other_qrf"),
    ("negative_city_signals", "BL Data 1", "city_name, state_name, location_pref, txn_30_days, txn_7_days"),
    ("churn_indicators", "pns_csd_daily_cron + Client Base", "pns_call, success_call, seller_rating, catalog_service, enabled"),
]
table5 = doc.add_table(rows=len(neg_keys) + 1, cols=3)
table5.style = "Light Grid Accent 1"
hdr5 = table5.rows[0].cells
hdr5[0].text = "Section"
hdr5[1].text = "Source Sheet(s)"
hdr5[2].text = "Keys Used"
for i, (section, source, keys) in enumerate(neg_keys):
    row = table5.rows[i + 1].cells
    row[0].text = section
    row[1].text = source
    row[2].text = keys

doc.add_heading("Full Prompt", level=2)
neg_prompt = """You are a Negative Engagement Analysis Agent for an enterprise seller platform.

Your job is to analyze the seller's data and identify ALL negative signals, risks, and friction points that indicate dissatisfaction, churn risk, or poor platform experience.

Do NOT decide whether to upsell.
Do NOT recommend any service.
Do NOT mention company name or personal identity.
Do NOT invent missing information.
Use "Unknown" when data is missing or unclear.

Focus on these signal categories:
1. NI (Not Interested) Signals: How many buyer enquiries did the seller reject? Break down by wrong category, location mismatch, retail, and other.
2. QRF (Quality Rejection Factor): Irrelevant or low-quality leads the seller received. Break down by wrong category, location mismatch, retail, and other.
3. Negative City Signals: Is there a geographic mismatch between the seller's location preference and the leads they are receiving?
4. Churn Indicators: Low PNS calls, low success calls, poor seller rating, or disabled account.

Seller Negative Engagement Data:
{seller_feature_json}

Return JSON in this schema:
{
  "ni_severity": "High / Medium / Low / None",
  "qrf_severity": "High / Medium / Low / None",
  "city_mismatch_risk": "High / Medium / Low / None",
  "churn_risk": "High / Medium / Low / None",
  "overall_negative_score": "Critical / Concerning / Low / None",
  "negative_signals": [],
  "evidence_fields": [],
  "negative_summary": ""
}"""
doc.add_paragraph(neg_prompt, style="No Spacing")

doc.add_page_break()

# ═══════════════════════════════════════════════════
# SECTION 6: AGENT 4 - Sentiment Analysis Agent
# ═══════════════════════════════════════════════════
doc.add_heading("6. Agent 4: Sentiment Analysis Agent", level=1)

doc.add_heading("Purpose", level=2)
doc.add_paragraph("Analyzes the seller's latest DSR conversation summary along with positive and negative engagement outputs to determine overall customer sentiment and upsell receptiveness.")

doc.add_heading("Input Data Keys", level=2)
sent_keys = [
    ("latest_dsr_summary", "DSR Summary", "summary, status, insertion_date, updated_on"),
    ("communication_context", "csd_kcd_active_seller + pns_csd_daily_cron + Scorecard", "totalcalls_30, totalcalls_60, total_ans_calls_30, pns_call, success_call, seller_rating, avg_ratings, connect_duration, last_call_date, dsr_meeting_last_date"),
    ("seller_context", "Client Base + Scorecard", "custtype_name, catalog_service, highest_service"),
    ("positive_engagement_analysis", "Output from Positive Agent", "bl_consumption_level, lms_engagement_level, enquiry_responsiveness, bl_activity_level, notification_engagement, overall_positive_score, positive_signals, positive_summary"),
    ("negative_engagement_analysis", "Output from Negative Agent", "ni_severity, qrf_severity, city_mismatch_risk, churn_risk, overall_negative_score, negative_signals, negative_summary"),
]
table6 = doc.add_table(rows=len(sent_keys) + 1, cols=3)
table6.style = "Light Grid Accent 1"
hdr6 = table6.rows[0].cells
hdr6[0].text = "Section"
hdr6[1].text = "Source"
hdr6[2].text = "Keys Used"
for i, (section, source, keys) in enumerate(sent_keys):
    row = table6.rows[i + 1].cells
    row[0].text = section
    row[1].text = source
    row[2].text = keys

doc.add_heading("Full Prompt", level=2)
sent_prompt = """You are a Sentiment Analysis Agent for an enterprise seller platform.

Your job is to analyze the seller's latest DSR (Daily Status Report) conversation summary along with their positive and negative engagement signals to determine the overall customer sentiment.

Do NOT decide whether to upsell.
Do NOT recommend any service.
Do NOT mention company name or personal identity.
Do NOT invent missing information.
Use "Unknown" when data is missing or unclear.

You will receive:
1. DSR Summary: The actual customer service conversation text and status.
2. Communication Context: Call volumes, ratings, connect duration.
3. Positive Engagement Analysis: Output from the Positive Agent (BL consumption, LMS, enquiries, notifications).
4. Negative Engagement Analysis: Output from the Negative Agent (NI rejections, QRF issues, churn risk).

Analyze all of this to determine:
1. Customer tone (frustrated, neutral, satisfied, happy)
2. Issue severity (were there critical complaints or minor queries?)
3. Resolution status (was the issue resolved or pending?)
4. Escalation risk (is the customer likely to escalate or churn?)
5. Upsell receptiveness (considering positive signals vs negative friction, is the seller open to upgrades?)

DSR, Communication & Agent Analysis Context:
{seller_feature_json}

Return JSON in this schema:
{
  "customer_tone": "Frustrated / Neutral / Satisfied / Happy / Unknown",
  "issue_severity": "Critical / Moderate / Minor / None / Unknown",
  "resolution_status": "Resolved / Pending / Escalated / Unknown",
  "escalation_risk": "High / Medium / Low / None",
  "upsell_receptiveness": "High / Medium / Low / None",
  "overall_sentiment": "Very Negative / Negative / Neutral / Positive / Very Positive / Unknown",
  "sentiment_signals": [],
  "evidence_fields": [],
  "sentiment_summary": ""
}"""
doc.add_paragraph(sent_prompt, style="No Spacing")

doc.add_page_break()

# ═══════════════════════════════════════════════════
# SECTION 7: AGENT 5 - Final Upsell Decision Agent
# ═══════════════════════════════════════════════════
doc.add_heading("7. Agent 5: Final Upsell Decision Agent", level=1)

doc.add_heading("Purpose", level=2)
doc.add_paragraph("The culmination agent that synthesizes ALL outputs from the 4 previous agents plus the prior segment recommendation to render the final upsell decision.")

doc.add_heading("Input Data Keys", level=2)
final_keys = [
    ("profile_status", "Output from Profile Agent", "seller_profile_type, business_maturity, client_vintage, current_customer_type, service_status, business_category_strength, catalog_completeness, engagement_level, location_profile, trust_profile, profile_risks, profile_summary"),
    ("positive_engagement", "Output from Positive Agent", "bl_consumption_level, lms_engagement_level, enquiry_responsiveness, bl_activity_level, notification_engagement, overall_positive_score, positive_signals, positive_summary"),
    ("negative_engagement", "Output from Negative Agent", "ni_severity, qrf_severity, city_mismatch_risk, churn_risk, overall_negative_score, negative_signals, negative_summary"),
    ("sentiment_analysis", "Output from Sentiment Agent", "customer_tone, issue_severity, resolution_status, escalation_risk, upsell_receptiveness, overall_sentiment, sentiment_signals, sentiment_summary"),
    ("prior_segment", "Upsell sheet (Master File)", "Segment Recommendation"),
]
table7 = doc.add_table(rows=len(final_keys) + 1, cols=3)
table7.style = "Light Grid Accent 1"
hdr7 = table7.rows[0].cells
hdr7[0].text = "Section"
hdr7[1].text = "Source"
hdr7[2].text = "Keys Used"
for i, (section, source, keys) in enumerate(final_keys):
    row = table7.rows[i + 1].cells
    row[0].text = section
    row[1].text = source
    row[2].text = keys

doc.add_heading("Full Prompt", level=2)
final_prompt = """You are the Final Upsell Decision Agent for an enterprise seller platform.

You are the LAST agent in a multi-agent pipeline. You have received the analyzed outputs from 4 specialized agents:
1. Profile Agent: Business profile, maturity, service status, category strength
2. Positive Engagement Agent: BL consumption, LMS usage, enquiry volume, notification engagement
3. Negative Engagement Agent: NI rejections, QRF quality issues, city mismatches, churn risk
4. Sentiment Agent: Customer tone, issue resolution, escalation risk

Your job is to synthesize ALL of these inputs and make a FINAL upsell recommendation.

Rules:
- Do NOT invent information not present in the agent outputs.
- If the Negative signals are Critical or Sentiment is Very Negative, recommend against upselling.
- If the Positive signals are Strong and Sentiment is Positive, strongly recommend upselling.
- Always provide a clear rationale backed by evidence from each agent.
- Recommend a specific service upgrade path based on the seller's current service level.

Service Hierarchy (lowest to highest):
MDC < MDC Pro < TrustSEAL < Maximiser < Star < Exporter

Agent Outputs:
{agent_outputs_json}

Prior Segment Recommendation (if available):
{prior_segment}

Return JSON in this schema:
{
  "upsell_decision": "Strongly Recommend / Recommend / Hold / Do Not Upsell",
  "confidence_score": 0.0,
  "recommended_service": "",
  "current_service": "",
  "upgrade_path": "",
  "key_positive_factors": [],
  "key_negative_factors": [],
  "key_sentiment_factors": [],
  "risk_assessment": "Low / Medium / High",
  "timing_recommendation": "Immediate / Next Renewal / After Issue Resolution / Not Recommended",
  "decision_rationale": "",
  "evidence_summary": {
    "from_profile": "",
    "from_positive": "",
    "from_negative": "",
    "from_sentiment": ""
  }
}"""
doc.add_paragraph(final_prompt, style="No Spacing")

# ── Save ──
output_path = r"c:\hackathon\hackathon 10X backend\outputs\Agent_Prompts_and_Data_Keys.docx"
doc.save(output_path)
print(f"Document saved to: {output_path}")
